import os
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

DEFAULT_RUN_PLAN_PATH = os.environ.get(
    "RUN_PLAN_FILE",
    str(Path(__file__).resolve().parent / "data" / "RunPlan.xlsx"),
)
if (
    not Path(DEFAULT_RUN_PLAN_PATH).exists()
    and Path(r"C:\Users\sanka\Desktop\Python_Playground\RunPlan.xlsx").exists()
):
    DEFAULT_RUN_PLAN_PATH = r"C:\Users\sanka\Desktop\Python_Playground\RunPlan.xlsx"


def read_excel_range(file_path, sheet_name, range_string):
    """
    Reads a specific range from a specific sheet.
    Example range_string: 'A2:B2'
    """
    # Split the range into start/end columns and rows
    start_cell, end_cell = range_string.split(":")

    # Extract columns (letters) and rows (numbers)
    start_col = "".join([c for c in start_cell if c.isalpha()])
    end_col = "".join([c for c in end_cell if c.isalpha()])

    start_row = int("".join([c for c in start_cell if c.isdigit()]))
    end_row = int("".join([c for c in end_cell if c.isdigit()]))

    # Calculate rows to skip and total rows to read
    # skiprows is 0-indexed, so row 2 means skipping 1 row
    skip = start_row - 1
    nrows = (end_row - start_row) + 1
    cols = f"{start_col}:{end_col}"

    # Read the separate sheet
    return pd.read_excel(
        file_path,
        sheet_name=sheet_name,
        skiprows=skip,
        nrows=nrows,
        usecols=cols,
        header=None,
    )


def get_plan_info(file_path, labels=None):
    """Return optional label rows from column A of the Plan sheet."""
    wanted = {" ".join(str(label).split()).lower() for label in (labels or (
        "Feed SL ID", "Feed Cut-In", "Feed molecular weight"
    ))}
    path = Path(file_path)
    if path.suffix.lower() != ".xlsx":
        return []
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        rows = []
        for row in workbook["Plan"].iter_rows(values_only=True):
            label = " ".join(str(row[0]).split()) if row and row[0] is not None else ""
            if label.lower() in wanted:
                rows.append({
                    "parameter": label,
                    "units": row[1] if len(row) > 1 and row[1] is not None else "",
                    "value": row[2] if len(row) > 2 and row[2] is not None else "",
                })
        return rows
    finally:
        workbook.close()


def load_run_plan_data(file_path=None):
    """
    Loads df_plan, df_Loading_plan, df_DOE, and df_synFeed/df_feed from the Excel workbook.
    """
    path = file_path or DEFAULT_RUN_PLAN_PATH
    if not Path(path).exists():
        raise FileNotFoundError(f"RunPlan Excel file not found: {path}")

    # Read "Plan" sheet (A2:B2)
    df_plan_local = read_excel_range(path, sheet_name="Plan", range_string="A2:B2")
    df_loading_local = read_excel_range(path, sheet_name="Plan", range_string="A4:j14")

    doe_usecols = "A:M"
    if Path(path).suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=True, read_only=True)
        try:
            headers = next(workbook["DOE"].iter_rows(min_row=10, max_row=10, values_only=True))
            last_header = max(
                (index for index, value in enumerate(headers) if value is not None and str(value).strip()),
                default=12,
            )
            doe_usecols = f"A:{get_column_letter(last_header + 1)}"
        finally:
            workbook.close()
    df_doe_local = pd.read_excel(
        path, sheet_name="DOE", skiprows=9, usecols=doe_usecols, header=None
    )
    filled_rows = df_doe_local.replace(r"^\s*$", pd.NA, regex=True).notna().any(axis=1)
    if filled_rows.any():
        df_doe_local = df_doe_local.iloc[: filled_rows[filled_rows].index[-1] + 1]
    if Path(path).suffix.lower() == ".xlsx":
        workbook = load_workbook(path, data_only=True)
        try:
            sheet = workbook["DOE"]
            for merged in sheet.merged_cells.ranges:
                value = sheet.cell(merged.min_row, merged.min_col).value
                if value is None:
                    continue
                for row in range(
                    max(merged.min_row, 11),
                    min(merged.max_row, len(df_doe_local) + 9) + 1,
                ):
                    for column in range(
                        merged.min_col, min(merged.max_col, df_doe_local.shape[1]) + 1
                    ):
                        df_doe_local.iat[row - 10, column - 1] = value
        finally:
            workbook.close()
    else:
        df_doe_local.iloc[:, 0:2] = df_doe_local.iloc[:, 0:2].ffill(axis=0)

    # Read "Feed" sheet (range A2:E25)
    df_feed_local = read_excel_range(path, sheet_name="Feed", range_string="A2:E25")

    return {
        "df_plan": df_plan_local,
        "df_Loading_plan": df_loading_local,
        "df_Plan": df_loading_local,
        "df_DOE": df_doe_local,
        "df_synFeed": df_feed_local,
        "df_feed": df_feed_local,
        "plan_info": get_plan_info(path),
    }


# Initialize default module-level DataFrames if default file exists
file_path = DEFAULT_RUN_PLAN_PATH
if Path(file_path).exists():
    _data = load_run_plan_data(file_path)
    df_plan = _data["df_plan"]
    df_Loading_plan = _data["df_Loading_plan"]
    df_Plan = _data["df_Plan"]
    df_DOE = _data["df_DOE"]
    df_synFeed = _data["df_synFeed"]
    df_feed = _data["df_feed"]
else:
    df_plan = pd.DataFrame()
    df_Loading_plan = pd.DataFrame()
    df_Plan = pd.DataFrame()
    df_DOE = pd.DataFrame()
    df_synFeed = pd.DataFrame()
    df_feed = pd.DataFrame()


if __name__ == "__main__":
    print("df_plan:")
    print(df_plan)
    print("\ndf_Plan (Loading Plan):")
    print(df_Plan)
    print("\ndf_DOE shape:", df_DOE.shape)
    print("\ndf_feed shape:", df_feed.shape)
