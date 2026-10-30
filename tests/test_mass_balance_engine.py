"""Golden tests: the Mass Balance engine must reproduce R6-C1-S1-MassBal_Ver2.2.xlsx (Excel-calculated values)."""
import json
from pathlib import Path

import pytest
from openpyxl.utils import get_column_letter

from backend.mass_balance.constants import MB_NAMES, ZMB_GASES
from backend.mass_balance.engine import calculate
from backend.mass_balance.service import zmb_gc_from_export
from backend.mass_balance.txt_loader import read_export
from backend.mass_balance.workbook import read_workbook

FIXTURES = Path(__file__).parent / "fixtures" / "mass_balance"
GOLDEN = json.loads((FIXTURES / "golden_R6-C1-S1.json").read_text())
TOL = 1e-7


@pytest.fixture(scope="module")
def result():
    data = read_workbook(FIXTURES / "R6-C1-S1-MassBal_Ver2.2.xlsx")
    return calculate(data["common"], data["samples"][0])


def same(mine, cell_value):
    if not isinstance(cell_value, (int, float)):
        return mine is None or mine == 0
    return mine is not None and abs(mine - cell_value) <= TOL * max(1.0, abs(cell_value))


def check(pairs):
    bad = [(label, mine, ref) for label, mine, ref in pairs if not same(mine, ref)]
    assert not bad, "\n".join(f"{label}: engine {mine} workbook {ref}" for label, mine, ref in bad[:20])
    return len(pairs)


def test_zmb_txt_equals_workbook_zmb_blocks():
    """The ZMB sheet's GC blocks are the ESTD values of each Sampled_reactor in the ZMB export."""
    export = read_export(FIXTURES / "Run6_ZMB_g_07_May_24_1.txt")
    blocks = zmb_gc_from_export(export)
    from backend.mass_balance.engine import zmb_stage
    zmb = zmb_stage({"gc": blocks})
    pairs = []
    for r, start in zip(range(1, 9), range(17, 153, 17)):
        for k, (name, _, _) in enumerate(ZMB_GASES):
            pairs.append((f"ZMB R{r} {name}", zmb["gc"][str(r)]["average"][name], GOLDEN["ZMB"][f"B{start + 5 + k}"]))
            pairs.append((f"ZMB R{r} {name} wt%", zmb["gc"][str(r)]["wt"][name], GOLDEN["ZMB"][f"C{start + 5 + k}"]))
    assert check(pairs) == 160


def test_zmb_liquid_and_run_plan(result):
    pairs = [(f"ZMB {get_column_letter(1 + r)}13", result["zmb"]["liquid"][str(r)]["verified"], GOLDEN["ZMB"][f"{get_column_letter(1 + r)}13"])
             for r in range(1, 9)]
    names = ["Hydrogen", "Ethylene", "Methane", "Ethane", "Nitrogen", "CO2", "CO", "Acetylene", "Propylene", "Propane"]
    for i, name in enumerate(names):
        col = get_column_letter(2 + i)
        pairs.append((f"Run_Plan {col}35", result["mixed_gas"]["mol"][name], GOLDEN["Run_Plan"][f"{col}35"]))
        pairs.append((f"Run_Plan {col}36", result["mixed_gas"]["mass"][name], GOLDEN["Run_Plan"][f"{col}36"]))
        pairs.append((f"Run_Plan {col}38", result["mixed_gas"]["g38"][name], GOLDEN["Run_Plan"][f"{col}38"]))
    pairs.append(("Run_Plan L35", result["mixed_gas"]["mol_sum"], GOLDEN["Run_Plan"]["L35"]))
    check(pairs)


def test_gas_sheet(result):
    g = result["gas"]
    pairs = [("Gas J36", g["mw_mix"], GOLDEN["Gas"]["J36"]), ("Gas I37", g["mass"]["Argon"], GOLDEN["Gas"]["I37"]),
             ("Gas J37", g["mass"]["Mixed gas"], GOLDEN["Gas"]["J37"])]
    for r in range(1, 9):
        col = get_column_letter(1 + r)
        pairs += [(f"Gas {col}12", g["reactors"][str(r)]["total"], GOLDEN["Gas"][f"{col}12"]),
                  (f"Gas {col}20", g["reactors"][str(r)]["argon"], GOLDEN["Gas"][f"{col}20"]),
                  (f"Gas {col}21", g["reactors"][str(r)]["contained"]["Mixed gas"], GOLDEN["Gas"][f"{col}21"])]
    check(pairs)


def test_liq_sheet(result):
    pairs = []
    for r in range(1, 9):
        if r == 3:          # no collection times entered: the workbook shows date arithmetic on blanks
            continue
        col = get_column_letter(1 + r)
        L = result["liq"][str(r)]
        pairs += [(f"Liq {col}7", L["duration_min"], GOLDEN["Liq"][f"{col}7"]), (f"Liq {col}8", L["tos_h"], GOLDEN["Liq"][f"{col}8"]),
                  (f"Liq {col}13", L["feed"], GOLDEN["Liq"][f"{col}13"]), (f"Liq {col}40", L["hc_flow"], GOLDEN["Liq"][f"{col}40"]),
                  (f"Liq {col}41", L["aq_flow"], GOLDEN["Liq"][f"{col}41"])]
    check(pairs)


@pytest.mark.parametrize("reactor", [1, 2, 4, 5, 6, 7, 8])
def test_reactor_sheets(result, reactor):
    R, gold = result["reactors"][str(reactor)], GOLDEN[f"R{reactor}_MB"]
    h = R["header"]
    pairs = [("F5", h["ar_in"], gold["F5"]), ("C8", h["liq_feed"], gold["C8"]), ("C9", h["gas_feed"], gold["C9"]),
             ("C10", h["hc_flow"], gold["C10"]), ("C11", h["aq_flow"], gold["C11"]), ("C12", h["gas_product"], gold["C12"]),
             ("C13", h["closure"], gold["C13"]), ("C86", R["mean_vtot"], gold["C86"]), ("C150", R["carbon_gas"], gold["C150"])]
    for k, name in enumerate(MB_NAMES):
        pairs.append((f"C{52 + k} {name}", R["mean_m"][name], gold[f"C{52 + k}"]))
        pairs.append((f"AC{52 + k} {name}", R["carbon_mean"][name], gold[f"AC{52 + k}"]))
    for j, rep in enumerate(R["reps"]):
        col = get_column_letter(4 + j)
        pairs.append((f"{col}150", rep["carbon_total"], gold[f"{col}150"]))
        pairs.append((f"{col}78", rep["m"]["Acetylene"], gold[f"{col}78"]))
    assert len(R["reps"]) == len(next(iter(gold["inputs_GC_estd_mol%"].values())))
    check([(f"R{reactor}_MB {label}", mine, ref) for label, mine, ref in pairs])


def test_c_balance(result):
    keys = {3: "c_liq_in", 4: "c_gas_in", 6: "c_liq_out", 7: "c_gas_out", 9: "carbon_balance", 11: "c4_liq_feed",
            12: "c4_liq_out", 14: "c4_gas_out", 15: "c4_yield", 17: "acetylene_in", 18: "acetylene_out",
            19: "acetylene_conversion", 20: "ethylene_ethane_ratio", 21: "ethylene_selectivity"}
    pairs = []
    for r in range(1, 9):
        if r == 3:
            continue
        col = get_column_letter(2 + r)
        C = result["c_balance"][str(r)]
        pairs += [(f"C-Balance {col}{row}", C[key], GOLDEN["C-Balance"][f"{col}{row}"]) for row, key in keys.items()]
        helper = get_column_letter(13 + r)
        hp = C["helpers"]
        for row, value in ((7, hp["Acetylene"]["c_flow"]), (17, hp["Methane"]["c_flow"]), (24, hp["Ethane"]["c_flow"]),
                           (31, hp["Ethylene"]["c_flow"]), (25, hp["Ethane"]["mass_flow"]), (32, hp["Ethylene"]["mass_flow"]),
                           (35, C["c_liq_out_g"]), (36, C["c_liq_out"])):
            pairs.append((f"C-Balance {helper}{row}", value, GOLDEN["C-Balance"][f"{helper}{row}"]))
    check(pairs)


def test_summary(result):
    s, gold = result["summary"], GOLDEN["Summary"]
    pairs = []
    for row in s["rows"]:
        if row["reactor"] == 3:
            continue
        r = 3 + row["reactor"]
        pairs += [(f"F{r}", row["carbon_balance"], gold[f"F{r}"]), (f"G{r}", row["mass_balance"], gold[f"G{r}"]),
                  (f"H{r}", row["tos_h"], gold[f"H{r}"]), (f"I{r}", row["acetylene_conversion"], gold[f"I{r}"]),
                  (f"J{r}", row["ethylene_selectivity"], gold[f"J{r}"]), (f"K{r}", row["carbon_gas"], gold[f"K{r}"]),
                  (f"L{r}", row["c4_yield"], gold[f"L{r}"])]
    for r in range(1, 9):
        for j in range(25):
            p = f"{get_column_letter(15 + r)}{4 + j}"
            a = f"{get_column_letter(5 + r)}{16 + j}"
            cb = s["carbon_balance_by_rep"][str(r)]
            ac = s["acetylene_conversion_by_rep"][str(r)]
            pairs.append((f"Summary {p}", cb[j] if j < len(cb) else None, gold[p]))
            pairs.append((f"Summary {a}", ac[j] if j < len(ac) else None, gold[a]))
    assert check(pairs) > 400


def test_corrected_mode_keeps_golden_values():
    data = read_workbook(FIXTURES / "R6-C1-S1-MassBal_Ver2.2.xlsx")
    res = calculate(data["common"], data["samples"][0], {"workbook_compatible": False})
    for r, expected in ((1, 107.36426096079839), (8, 86.73060641620319)):
        assert same(res["c_balance"][str(r)]["carbon_balance"], expected)
    assert same(res["c_balance"]["1"]["ethylene_selectivity"], GOLDEN["C-Balance"]["C21"])

