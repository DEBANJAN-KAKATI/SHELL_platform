"""Reforming API (/api/reforming/*). Isomerization routes stay in backend/routes/api.py."""

from pathlib import Path

from flask import Blueprint, jsonify, request

from backend.reforming.operating import read_operating_file
from backend.services.run_plan_service import process_run_plan_file
from backend.reforming.service import process_reforming_directory
from backend.reforming.sync_source import save_folder, sync_status

BASE = Path(__file__).resolve().parents[2]
reforming_api = Blueprint("reforming_api", __name__, url_prefix="/api/reforming")


@reforming_api.post("/process")
def process():
    folder = (request.get_json(silent=True) or {}).get("path")
    if not folder:
        return jsonify({"error": "Folder path is required."}), 400
    try:
        return jsonify(process_reforming_directory(folder))
    except Exception as error:
        return jsonify({"error": str(error)}), 400


@reforming_api.get("/settings")
def settings():
    from flask import current_app
    path = current_app.config.get("DEFAULT_RUN_PLAN_FILE", "")
    return jsonify({"default_run_plan_file": path})


@reforming_api.get("/sync-source")
def sync_source():
    try:
        return jsonify(sync_status())
    except Exception as error:
        return jsonify({"found": False, "message": str(error)})


@reforming_api.post("/sync-source")
def set_sync_source():
    folder = (request.get_json(silent=True) or {}).get("folder", "")
    try:
        save_folder(folder)
        return jsonify(sync_status())
    except Exception as error:
        return jsonify({"error": str(error)}), 400


@reforming_api.post("/process-run-plan")
def process_run_plan():
    file_path = (request.get_json(silent=True) or {}).get("path")
    if not file_path:
        return jsonify({"error": "Run Plan file path is required."}), 400
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(file_path, read_only=True, data_only=True)
        try:
            plan_values = [
                row[0] for row in workbook["Plan"].iter_rows(
                    min_col=1, max_col=1, values_only=True
                ) if row[0] is not None
            ]
            if "Feed DHA" in workbook.sheetnames or any(
                str(value).strip().lower() == "run information" for value in plan_values
            ):
                return jsonify({
                    "error": "This is the legacy reforming Run Plan format. Use the shared RunPlan.xlsx template with Plan, DOE and Feed sheets."
                }), 400
        finally:
            workbook.close()
        return jsonify(process_run_plan_file(file_path))
    except Exception as error:
        return jsonify({"error": str(error)}), 400


@reforming_api.post("/operating-data")
def operating_data():
    file_path = (request.get_json(silent=True) or {}).get("path")
    if not file_path:
        return jsonify({"error": "Load the operating data (.txt) on the Online Analysis page first."}), 400
    try:
        return jsonify(read_operating_file(file_path))
    except Exception as error:
        return jsonify({"error": str(error)}), 400
