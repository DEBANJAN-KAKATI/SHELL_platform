"""Mass Balance Run Plan in the shared RunPlan.xlsx format (Plan, DOE, Feed) plus a Liq sheet.

- Plan: objective (B2), Catalyst Loading Plan (A4:J14, reactors 1-8 in C:J), Op.Conditions, and label rows in
  column A with the value in column C: Run Start (or Feed Cut-In), Feed SL ID, Feed molecular weight.
- DOE: header row 10, data from row 11 (one row per condition, block and reactor). Gas Flow (NmL/h) per reactor is the
  hand-entered feed flow used when no Online Analysis export is loaded; Temperature / Pressure per block.
- Feed: liquid feed (Carbon E2, Hydrogen E3, Molecular Weight E4, Oxygen, Sulfur, Nitrogen, Density, ID, description)
  and gas feed (verified Argon, verified mixed gas, total gas feed 01i14FI01, target mixed-gas composition), labels in
  column A, unit in D, value in E.
- Liq: "Liquid samples" table (one row per sample and reactor) and "ZMB liquid verification" table (3 reps).
The result has the same shape as workbook.read_workbook(): {file, path, common, samples}.
"""
import re
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

from .constants import REACTORS

MIX_NAMES = ["Hydrogen", "Ethylene", "Methane", "Ethane", "Nitrogen", "CO2", "CO", "Acetylene", "Propylene", "Propane"]


def is_run_plan_workbook(sheetnames):
    return {"Plan", "DOE", "Feed"} <= set(sheetnames) and "Run_Plan" not in sheetnames


def _norm(text):
    return re.sub(r"[^a-z0-9]+", " ", str(text or "").lower()).strip()


def _value(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, str):
        value = value.strip()
        if not value or value.upper() == "X" or value == "--":
            return None
        try:
            return float(value)
        except ValueError:
            return value
    return value


def _number(value):
    value = _value(value)
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _text(value):
    value = _value(value)
    return None if value is None else (str(int(value)) if isinstance(value, float) and value.is_integer() else str(value))


def _reactor(value):
    match = re.search(r"([1-8])", str(value or ""))
    return int(match.group(1)) if match else None


def _sheet_rows(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def _merged_fill(ws, rows):
    """Copy the value of each merged range into all its cells (DOE condition / block / dates)."""
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for r in range(merged.min_row, merged.max_row + 1):
            for col in range(merged.min_col, merged.max_col + 1):
                if r - 1 < len(rows) and col - 1 < len(rows[r - 1]):
                    rows[r - 1][col - 1] = value
    return rows


def _table(rows, header_index, aliases):
    """Rows below a header row until the first empty row, as dicts keyed by the alias names."""
    header = [_norm(h) for h in rows[header_index]]
    columns = {}
    for key, names in aliases.items():
        for i, h in enumerate(header):
            if any(h.startswith(_norm(n)) for n in names):
                columns[key] = i
                break
    out = []
    for row in rows[header_index + 1:]:
        if not any(v not in (None, "") for v in row):
            break
        out.append({key: row[i] if i < len(row) else None for key, i in columns.items()})
    return out


def read_plan(ws):
    rows = _sheet_rows(ws)
    labels = {}
    for row in rows:
        if row and row[0] is not None:
            labels[_norm(row[0])] = row
    objective = _text(labels.get("objective", [None, None])[1])
    catalysts, conditions = [], []
    reactor_row = next((i for i, row in enumerate(rows) if row and _norm(row[0]) == "reactor"), None)
    loading = {}
    if reactor_row is not None:
        for row in rows[reactor_row + 1:reactor_row + 12]:
            if not row or row[0] is None:
                continue
            loading[_norm(row[0])] = [row[2 + k] if 2 + k < len(row) else None for k in range(8)]
    get = lambda *names: next((loading[n] for n in loading if any(n.startswith(_norm(x)) for x in names)), [None] * 8)
    cat, volume, cbd = get("catalyst"), get("catalyst volume"), get("catalyst cbd")
    amount, loaded, shape = get("catalyst amount"), get("loaded catalyst"), get("shape")
    for k, r in enumerate(REACTORS):
        vol, dens = _number(volume[k]), _number(cbd[k])
        grams = _number(amount[k]) if _number(amount[k]) is not None else (vol * dens if vol is not None and dens is not None else None)
        beds = [{"bed": "Top", "catalyst": _text(cat[k]), "target_mg": grams * 1000 if grams is not None else None,
                 "actual_mg": _number(loaded[k]) * 1000 if _number(loaded[k]) is not None else None, "volume_ml": vol,
                 "shape": _text(shape[k])}]
        beds += [{"bed": b, "catalyst": None, "target_mg": None, "actual_mg": None, "volume_ml": None} for b in ("Middle", "Bottom")]
        catalysts.append({"reactor": r, "beds": beds})
    op_row = next((i for i, row in enumerate(rows) if row and _norm(row[0]).startswith("op conditions")), None)
    op = {}
    if op_row is not None:
        for row in rows[op_row + 1:op_row + 8]:
            if row and row[0] is not None:
                op[_norm(row[0])] = [row[2 + k] if 2 + k < len(row) else None for k in range(6)]
    for k in range(6):
        conditions.append({"condition": k + 1, "temperature_1": _number((op.get("temperature") or [None] * 6)[k]),
                           "pressure_1": _number((op.get("pressure") or [None] * 6)[k]),
                           "duration": _value((op.get("duration") or [None] * 6)[k]), "duration_unit": "hours"})
    info = {key: (_value(row[2]) if len(row) > 2 else None) for key, row in labels.items()}
    start = info.get("run start") or info.get("feed cut in")
    return {"objective": objective, "catalysts": catalysts, "conditions": conditions, "start": start,
            "run_number": _text(info.get("run number")), "feed_id": _text(info.get("feed sl id")),
            "feed_mw": _number(info.get("feed molecular weight"))}


def read_doe(ws):
    rows = _merged_fill(ws, _sheet_rows(ws))
    header = next((i for i, row in enumerate(rows) if row and _norm(row[0]).startswith("condition")), 9)
    table = _table(rows, header, {
        "condition": ["condition"], "block": ["block"], "start": ["start date"], "end": ["end date"],
        "temperature": ["temperature"], "pressure": ["pressure"], "reactor": ["reactor"], "gas_flow": ["gas flow"],
        "liquid_flow": ["liquid flow"], "lhsv": ["lhsv"], "gor": ["gor"], "feed_gas": ["feed gas"], "remarks": ["remarks"],
        "catalyst": ["catalyst"]})
    doe = []
    for row in table:
        r = _reactor(row.get("reactor"))
        if not r:
            continue
        doe.append({"condition": _text(row.get("condition")), "block": _text(row.get("block")), "start": _value(row.get("start")),
                    "end": _value(row.get("end")), "temperature": _number(row.get("temperature")), "pressure": _number(row.get("pressure")),
                    "reactor": r, "gas_flow": _number(row.get("gas_flow")), "liquid_flow": _number(row.get("liquid_flow")),
                    "lhsv": _number(row.get("lhsv")), "gor": _number(row.get("gor")), "remarks": _text(row.get("remarks")),
                    "catalyst": _text(row.get("catalyst"))})
    return doe


def read_feed(ws):
    values = {}
    for row in ws.iter_rows(min_row=1, max_row=40, values_only=True):
        if row and row[0] is not None:
            values[_norm(row[0])] = row[4] if len(row) > 4 else None
    find = lambda *names: next((values[k] for k in values if any(k.startswith(_norm(n)) for n in names)), None)
    mixed = {name: _number(find(f"mixed gas {name}")) for name in MIX_NAMES}
    return {
        "liquid_feed_carbon_wt": _number(find("carbon")), "liquid_feed_hydrogen_wt": _number(find("hydrogen")),
        "liquid_feed_mw": _number(find("molecular weight")), "liquid_feed_oxygen_wt": _number(find("oxygen")),
        "liquid_feed_sulfur_ppm": _number(find("sulfur")), "liquid_feed_nitrogen_ppm": _number(find("nitrogen")),
        "liquid_feed_density": _number(find("density")), "liquid_feed_id": _text(find("feed id")),
        "liquid_feed_description": _text(find("feed description")),
        "ar_mol": _number(find("argon")), "mix_mol": _number(find("mixed gas verified")),
        "fi01": _number(find("total gas feed")),
        "mixed_gas_target": mixed if any(v is not None for v in mixed.values()) else None,
    }


def read_liq(ws):
    rows = _sheet_rows(ws)
    samples, zmb = [], {}
    for i, row in enumerate(rows):
        first = _norm(row[0]) if row else ""
        if first == "sample":
            for item in _table(rows, i, {"sample": ["sample"], "condition": ["condition"], "reactor": ["reactor"],
                                         "start": ["collection started"], "finish": ["collection finished"], "tlp": ["tlp"],
                                         "hc": ["hc", "hydrocarbon"], "aq": ["aqueous"], "c_wt": ["c in liquid", "carbon in liquid"],
                                         "density": ["density"]}):
                if item.get("sample") and _reactor(item.get("reactor")):
                    samples.append(item)
        if first == "rep":
            for item in _table(rows, i, {"rep": ["rep"], "reactor": ["reactor"], "start": ["collection started"],
                                         "finish": ["collection finished"], "mass": ["tl collected", "collected mass"]}):
                rep, r = _text(item.get("rep")), _reactor(item.get("reactor"))
                if rep and r:
                    rep = re.sub(r"\D", "", rep) or rep
                    zmb.setdefault(rep, {})[str(r)] = {"start": _value(item.get("start")), "finish": _value(item.get("finish")),
                                                       "mass": _number(item.get("mass"))}
    return samples, zmb


def read_run_plan(path):
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"File not found: {path}")
    wb = load_workbook(path, data_only=True)
    if not is_run_plan_workbook(wb.sheetnames):
        raise ValueError(f"{path.name} is not a Run Plan workbook (sheets Plan, DOE and Feed are needed).")
    plan, doe, feed = read_plan(wb["Plan"]), read_doe(wb["DOE"]), read_feed(wb["Feed"])
    liq_rows, zmb_liquid = read_liq(wb["Liq"]) if "Liq" in wb.sheetnames else ([], {})
    wb.close()

    # Conditions: temperature / pressure of block 1 (R1-R4) and block 2 (R5-R8) from the DOE, else Plan Op.Conditions.
    conditions = plan["conditions"]
    for cond in conditions:
        rows = [d for d in doe if d["condition"] == str(cond["condition"])]
        b1 = next((d for d in rows if d["reactor"] <= 4 and d["temperature"] is not None), None)
        b2 = next((d for d in rows if d["reactor"] >= 5 and d["temperature"] is not None), None)
        if b1:
            cond.update(temperature_1=b1["temperature"], pressure_1=b1["pressure"])
        if b2:
            cond.update(temperature_2=b2["temperature"], pressure_2=b2["pressure"])
        dated = next((d for d in rows if d["start"] and d["end"]), None)
        if dated:
            hours = (datetime.fromisoformat(dated["end"]) - datetime.fromisoformat(dated["start"])).total_seconds() / 3600
            cond.update(duration=round(hours, 2), duration_unit="hours")
    for d in doe:                                   # DOE catalyst column overrides the Plan catalyst of that reactor
        if d["catalyst"]:
            plan["catalysts"][d["reactor"] - 1]["beds"][0]["catalyst"] = d["catalyst"]
    start = plan["start"] or min((d["start"] for d in doe if d["start"]), default=None)

    run_plan = {key: feed[key] for key in feed if key.startswith("liquid_feed")}
    run_plan["liquid_feed_id"] = run_plan.get("liquid_feed_id") or plan["feed_id"]
    run_plan.update(catalysts=plan["catalysts"], conditions=conditions, doe=doe, mixed_gas_target=feed["mixed_gas_target"])
    common = {"run": {"run_number": plan["run_number"], "objective": plan["objective"], "start": start},
              "run_plan": run_plan, "zmb": {"liquid": zmb_liquid},
              "gas": {"pure_mol": {}, "ar_mol": feed["ar_mol"], "mix_mol": feed["mix_mol"]}}

    samples = []
    for item in liq_rows:
        name = _text(item["sample"])
        sample = next((s for s in samples if s["name"] == name), None)
        if sample is None:
            cond = _text(item.get("condition")) or (re.search(r"C(\d+)", name, re.I).group(1) if re.search(r"C(\d+)", name, re.I) else None)
            nml = {str(d["reactor"]): d["gas_flow"] for d in doe if cond and d["condition"] == cond}
            sample = {"name": name, "condition": cond, "liq": {}, "c_wt_liquid_product": {},
                      "gas": {"source": "txt", "fi01": feed["fi01"], "nml": {str(r): nml.get(str(r)) or 0 for r in REACTORS}},
                      "gc_source": "txt", "gc": {}, "excluded": {}, "window": {}}
            samples.append(sample)
        r = str(_reactor(item["reactor"]))
        sample["liq"][r] = {"start": _value(item.get("start")), "finish": _value(item.get("finish")), "tlp": _number(item.get("tlp")),
                            "hc": _number(item.get("hc")), "aq": _number(item.get("aq")),
                            "analyses": {"Density [g/mL]": _number(item.get("density"))}}
        if _number(item.get("c_wt")) is not None:
            sample["c_wt_liquid_product"][r] = _number(item.get("c_wt"))
    if not samples:
        samples = [{"name": "C1-S1", "liq": {}, "c_wt_liquid_product": {}, "gas": {"source": "txt", "fi01": feed["fi01"], "nml": {}},
                    "gc_source": "txt", "gc": {}, "excluded": {}, "window": {}}]
    return {"file": path.name, "path": str(path), "format": "run_plan", "common": common, "samples": samples}
