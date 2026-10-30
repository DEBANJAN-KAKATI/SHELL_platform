"""Mass Balance API (/api/mass-balance/*)."""
from datetime import datetime

from flask import Blueprint, jsonify, request, send_file

from backend.services.dialog_service import browse_file_dialog
from .export import workbook_bytes
from .service import export_summary, load_export, read_gc_map, run_calculation, save_gc_map
from .workbook import read_workbook
from .run_plan import is_run_plan_workbook, read_run_plan
from .constants import MB_NAMES

mass_balance_api = Blueprint("mass_balance_api", __name__, url_prefix="/api/mass-balance")


def _error(message, status=400):
    return jsonify({"error": message}), status


@mass_balance_api.post("/browse-file")
def browse_file():
    body = request.get_json(silent=True) or {}
    kind = body.get("kind", "txt")
    filetypes = [("Excel workbooks", "*.xlsx"), ("All files", "*.*")] if kind == "xlsx" else [("Text files", "*.txt"), ("All files", "*.*")]
    try:
        selected, cancelled = browse_file_dialog(initial_file=body.get("initial_file") or "", title=body.get("title") or "Select file",
                                                 filetypes=filetypes)
        return jsonify({"path": selected, "cancelled": cancelled})
    except Exception as error:
        return jsonify({"error": "Native file browser is unavailable. Enter the file path manually.", "details": str(error)}), 503


@mass_balance_api.post("/import-txt")
def import_txt():
    path = (request.get_json(silent=True) or {}).get("path")
    if not path:
        return _error("File path is required.")
    try:
        return jsonify(export_summary(load_export(path)))
    except Exception as error:
        return _error(str(error))


@mass_balance_api.post("/import-workbook")
def import_workbook():
    path = (request.get_json(silent=True) or {}).get("path")
    if not path:
        return _error("Workbook path is required.")
    try:
        from openpyxl import load_workbook
        workbook = load_workbook(path, read_only=True)
        sheets = workbook.sheetnames
        workbook.close()
        if is_run_plan_workbook(sheets):               # RunPlan.xlsx format (Plan, DOE, Feed, Liq)
            return jsonify(read_run_plan(path))
        return jsonify(read_workbook(path))            # filled MassBal workbook (Run_Plan, Liq, Gas, ZMB, ...)
    except Exception as error:
        return _error(str(error))


@mass_balance_api.get("/gc-map")
def get_gc_map():
    return jsonify({"map": read_gc_map(), "rows": MB_NAMES})


@mass_balance_api.post("/gc-map")
def set_gc_map():
    try:
        return jsonify({"map": save_gc_map((request.get_json(silent=True) or {}).get("map") or {}), "rows": MB_NAMES})
    except Exception as error:
        return _error(str(error))


@mass_balance_api.post("/calculate")
def calculate():
    try:
        return jsonify(run_calculation(request.get_json(silent=True) or {}))
    except Exception as error:
        return _error(str(error))


@mass_balance_api.post("/export-xlsx")
def export_xlsx():
    try:
        out = run_calculation(request.get_json(silent=True) or {})
        run = ((request.get_json(silent=True) or {}).get("common") or {}).get("run") or {}
        name = f"MassBalance_{run.get('run_number') or 'run'}_{datetime.now():%Y%m%d_%H%M}.xlsx".replace(" ", "_")
        return send_file(workbook_bytes(out["results"]), as_attachment=True, download_name=name,
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as error:
        return _error(str(error))
