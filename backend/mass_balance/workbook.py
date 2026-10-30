"""Read the hand-entered (blue) cells of a MassBal workbook (Ver2.x) into the engine's input format.

A filled workbook can be loaded as it is: its run/plan values, ZMB liquid reps, verified feed gas, Liq sample, the
carbon wt% of the liquid product and, when no .txt is imported, the GC tables of the ZMB and Rx_MB sheets.
"""
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .constants import GAS_PURE, MB_NAMES, REACTORS, ZMB_GASES

ZMB_BLOCK_START = {str(r): 17 + 17 * (r - 1) for r in REACTORS}
ZMB_BLOCK_START["0"] = 152                     # "GC Analysis Rep Number for Bypass"
REP_COLUMNS = range(4, 29)                     # D:AB
LIQ_REP_COLUMNS = {"1": 2, "2": 12, "3": 22}   # B, L, V (R1 column of each ZMB liquid rep)


def _value(value):
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, str):
        value = value.strip()
        return value or None
    return value


def _number(value):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def read_workbook(path):
    path = Path(path)
    if not path.is_file():
        raise ValueError(f"File not found: {path}")
    wb = load_workbook(path, data_only=True)
    missing = [name for name in ("Run_Plan", "Liq", "Gas", "ZMB") if name not in wb.sheetnames]
    if missing:
        raise ValueError(f"{path.name} is not a MassBal workbook (missing sheets: {', '.join(missing)}).")
    rp, liq, gas, zmb = wb["Run_Plan"], wb["Liq"], wb["Gas"], wb["ZMB"]

    run = {"run_number": _value(rp["B1"].value), "objective": _value(rp["B2"].value), "start": _value(rp["B3"].value),
           "notebook": _value(rp["B4"].value)}
    catalysts = []
    for block_row, first_reactor in ((8, 1), (15, 5)):
        for i in range(4):
            col = 2 + 4 * i
            beds = []
            for bed, row in (("Top", block_row + 2), ("Middle", block_row + 3), ("Bottom", block_row + 4)):
                beds.append({"bed": bed, "catalyst": _value(rp.cell(row, col).value),
                             "target_mg": _value(rp.cell(row, col + 1).value), "actual_mg": _value(rp.cell(row, col + 2).value),
                             "volume_ml": _value(rp.cell(row, col + 3).value)})
            catalysts.append({"reactor": first_reactor + i, "beds": beds})
    conditions = []
    for i in range(6):
        col = 2 + 2 * i
        conditions.append({"condition": i + 1, "temperature_1": _value(rp.cell(81, col).value),
                           "temperature_2": _value(rp.cell(82, col).value), "pressure_1": _value(rp.cell(83, col).value),
                           "pressure_2": _value(rp.cell(84, col).value), "duration": _value(rp.cell(85, col).value),
                           "duration_unit": _value(rp.cell(85, col + 1).value)})
    run_plan = {"liquid_feed_id": _value(rp["B48"].value), "liquid_feed_description": _value(rp["B49"].value),
                "liquid_feed_density": _value(rp["B52"].value),
                "liquid_feed_carbon_wt": _value(rp["B58"].value), "liquid_feed_hydrogen_wt": _value(rp["B59"].value),
                "liquid_feed_oxygen_wt": _value(rp["B60"].value), "liquid_feed_sulfur_ppm": _value(rp["B61"].value),
                "liquid_feed_nitrogen_ppm": _value(rp["B62"].value), "catalysts": catalysts, "conditions": conditions}

    liquid = {}
    for rep, first_col in LIQ_REP_COLUMNS.items():
        liquid[rep] = {str(r): {"start": _value(zmb.cell(4, first_col + r - 1).value),
                                "finish": _value(zmb.cell(5, first_col + r - 1).value),
                                "mass": _number(zmb.cell(10, first_col + r - 1).value)} for r in REACTORS}
    zmb_gc = {}
    for stream, start in ZMB_BLOCK_START.items():
        injections = []
        for col in REP_COLUMNS:
            values = {name: _number(zmb.cell(start + 5 + k, col).value) for k, (name, _, _) in enumerate(ZMB_GASES)}
            if any(v is not None for v in values.values()):
                injections.append({"time": _value(zmb.cell(start + 1, col).value), "estd": values,
                                   "source": f"workbook ZMB!{get_column_letter(col)}{start + 5}"})
        if injections:
            zmb_gc[stream] = injections

    gas_in = {"pure_mol": {name: _number(gas.cell(35, 2 + i).value) for i, (name, _) in enumerate(GAS_PURE)},
              "ar_mol": _number(gas["I35"].value), "mix_mol": _number(gas["J35"].value)}
    sample_gas = {"fi01": _number(gas["E5"].value), "nml": {str(r): _number(gas.cell(11, 1 + r).value) for r in REACTORS},
                  "source": "manual"}

    sample_liq = {}
    for r in REACTORS:
        col = 1 + r
        sample_liq[str(r)] = {
            "start": _value(liq.cell(5, col).value), "finish": _value(liq.cell(6, col).value),
            "tlp": _number(liq.cell(17, col).value), "hc": _number(liq.cell(18, col).value), "aq": _number(liq.cell(19, col).value),
            "analyses": {label: _number(liq.cell(row, col).value) for row, label in
                         ((24, "Carbon [wt%]"), (25, "Hydrogen [wt%]"), (26, "Oxygen [wt%]"), (27, "Nitrogen [ppmw]"),
                          (28, "Sulfur [ppmw]"), (32, "Density [g/mL]"))},
        }
    c_wt = {}
    if "C-Balance" in wb.sheetnames:
        cb = wb["C-Balance"]
        c_wt = {str(r): _number(cb.cell(34, 13 + r).value) for r in REACTORS}
    gc = {}
    for r in REACTORS:
        name = f"R{r}_MB"
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        reps = []
        for col in REP_COLUMNS:
            y = {row_name: _number(ws.cell(20 + k, col).value) for k, row_name in enumerate(MB_NAMES)}
            if any(v for v in y.values()):
                reps.append({"time": _value(ws.cell(16, col).value), "y": y,
                             "source": f"workbook {name}!{get_column_letter(col)}20:{get_column_letter(col)}48"})
        gc[str(r)] = reps
    wb.close()
    sample = {"name": _value(liq["B4"].value) or "S1", "liq": sample_liq, "c_wt_liquid_product": c_wt, "gas": sample_gas,
              "gc": gc, "gc_source": "workbook"}
    return {"file": path.name, "path": str(path),
            "common": {"run": run, "run_plan": run_plan, "zmb": {"liquid": liquid, "gc": zmb_gc}, "gas": gas_in},
            "samples": [sample]}
