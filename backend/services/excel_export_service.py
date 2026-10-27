"""Builds a multi-sheet .xlsx from table data sent by the frontend (Yield export + Online Analysis tab)."""
import math
import re
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

_INVALID_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="071A2D")
MAX_ROWS = 200000
MAX_COLUMNS = 500


def _sheet_title(name, used):
    title = _INVALID_SHEET_CHARS.sub(" ", str(name or "Sheet")).strip()[:31] or "Sheet"
    base, n = title, 2
    while title.lower() in used:
        suffix = f" ({n})"
        title = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (int, float)):
        return value if math.isfinite(value) else None
    return str(value)


def build_workbook(sheets):
    """sheets: [{name, columns: [str], rows: [[value...]], number_formats: {column: format}}]"""
    if not isinstance(sheets, list) or not sheets:
        raise ValueError("No sheets to export.")
    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for sheet in sheets:
        columns = [str(c) for c in (sheet.get("columns") or [])][:MAX_COLUMNS]
        rows = sheet.get("rows") or []
        if len(rows) > MAX_ROWS:
            raise ValueError(f"Sheet '{sheet.get('name')}' has too many rows ({len(rows)}).")
        formats = sheet.get("number_formats") or {}
        ws = wb.create_sheet(_sheet_title(sheet.get("name"), used))
        ws.append(columns)
        for cell in ws[1]:
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
        widths = [len(c) for c in columns]
        for row in rows:
            values = [_cell_value(v) for v in list(row)[: len(columns)]]
            ws.append(values)
            for i, v in enumerate(values):
                if v is not None:
                    widths[i] = max(widths[i], len(str(v)) if isinstance(v, str) else 10)
        for i, column in enumerate(columns, start=1):
            letter = get_column_letter(i)
            ws.column_dimensions[letter].width = min(max(widths[i - 1] + 2, 8), 45)
            fmt = formats.get(column)
            if fmt and rows:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = fmt
        ws.freeze_panes = "B2" if columns else None
        if columns and rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
